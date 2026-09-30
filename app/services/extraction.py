"""Turn uploaded files into clean text segments.

Each segment carries the page (PDF) or slide number (PPTX) and the nearest
section heading, which later become the citations students see in answers.
"""
import io
import logging
import re
from collections import Counter
from dataclasses import dataclass

from ..utils.text import normalize_text

log = logging.getLogger(__name__)


class ExtractionError(Exception):
    pass


@dataclass
class Segment:
    text: str
    page: int | None = None
    section: str | None = None


def extract(data, extension):
    extension = extension.lower().lstrip(".")
    handlers = {
        "pdf": _extract_pdf,
        "docx": _extract_docx,
        "pptx": _extract_pptx,
        "txt": _extract_plain,
        "md": _extract_markdown,
    }
    handler = handlers.get(extension)
    if handler is None:
        raise ExtractionError(f"Unsupported file type: .{extension}")
    try:
        segments = handler(data)
    except ExtractionError:
        raise
    except Exception as exc:
        log.exception("extraction failed")
        raise ExtractionError(f"Could not read the file: {exc}") from exc

    segments = [s for s in segments if s.text and s.text.strip()]
    if not segments:
        raise ExtractionError(
            "No readable text found. If this is a scanned PDF, run it through OCR first."
        )
    return segments


def _looks_like_heading(line):
    line = line.strip()
    if not line or len(line) > 90 or line.endswith((".", ",", ";", ":")):
        return False
    words = line.split()
    if len(words) > 12:
        return False
    if re.match(r"^(\d+(\.\d+)*|[IVX]+\.|chapter|section|unit|module|lecture)\b", line, re.I):
        return True
    caps = sum(1 for w in words if w[:1].isupper())
    return caps >= max(1, int(len(words) * 0.7)) and len(words) <= 8


def strip_repeated_lines(pages, threshold=0.5):
    """Drop lines that show up on most pages, which are nearly always headers,
    footers, course codes or slide templates."""
    if len(pages) < 4:
        return pages
    counts = Counter()
    for text in pages:
        seen = {ln.strip() for ln in text.splitlines() if ln.strip()}
        counts.update(seen)
    limit = max(3, int(len(pages) * threshold))
    repeated = {ln for ln, c in counts.items() if c >= limit and len(ln) < 120}
    page_number = re.compile(r"^(page\s*)?\d+(\s*(of|/)\s*\d+)?$", re.I)
    cleaned = []
    for text in pages:
        lines = [
            ln for ln in text.splitlines()
            if ln.strip() not in repeated and not page_number.match(ln.strip())
        ]
        cleaned.append("\n".join(lines))
    return cleaned


def _extract_pdf(data):
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception:
            raise ExtractionError("This PDF is password protected")
    raw_pages = []
    for page in reader.pages:
        try:
            raw_pages.append(page.extract_text() or "")
        except Exception as exc:
            log.warning("skipping unreadable pdf page: %s", exc)
            raw_pages.append("")
    pages = strip_repeated_lines(raw_pages)

    segments = []
    section = None
    for number, text in enumerate(pages, start=1):
        text = normalize_text(text)
        if not text:
            continue
        first = text.splitlines()[0]
        if _looks_like_heading(first):
            section = first.strip()
        segments.append(Segment(text=text, page=number, section=section))
    return segments


def _extract_docx(data):
    import docx

    document = docx.Document(io.BytesIO(data))
    segments = []
    section = None
    buffer = []

    def flush():
        if buffer:
            segments.append(Segment(text=normalize_text("\n".join(buffer)), section=section))
            buffer.clear()

    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        style = (para.style.name or "").lower() if para.style is not None else ""
        if style.startswith("heading") or style == "title":
            flush()
            section = text
            continue
        buffer.append(text)
    flush()

    for table in document.tables:
        rows = []
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                rows.append(" | ".join(cells))
        if rows:
            segments.append(Segment(text="\n".join(rows), section=section or "Table"))
    return segments


def _extract_pptx(data):
    try:
        from pptx import Presentation
    except ImportError:
        raise ExtractionError("PowerPoint support needs python-pptx installed")
    deck = Presentation(io.BytesIO(data))
    segments = []
    for number, slide in enumerate(deck.slides, start=1):
        title = None
        lines = []
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            text = "\n".join(p.text for p in shape.text_frame.paragraphs if p.text.strip())
            if not text:
                continue
            if shape == getattr(slide.shapes, "title", None):
                title = text.strip()
            else:
                lines.append(text)
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                lines.append("Notes: " + notes)
        body = "\n".join(([title] if title else []) + lines)
        if body.strip():
            segments.append(Segment(text=normalize_text(body), page=number, section=title))
    return segments


def _decode(data):
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ExtractionError("Could not decode the text file")


def _extract_plain(data):
    text = normalize_text(_decode(data))
    segments = []
    section = None
    for block in re.split(r"\n\s*\n", text):
        block = block.strip()
        if not block:
            continue
        lines = block.splitlines()
        if len(lines) == 1 and _looks_like_heading(lines[0]):
            section = lines[0].strip()
            continue
        segments.append(Segment(text=block, section=section))
    return segments


def _extract_markdown(data):
    text = _decode(data)
    text = re.sub(r"```.*?```", lambda m: m.group(0).replace("\n\n", "\n"), text, flags=re.S)
    segments = []
    section = None
    buffer = []

    def flush():
        body = normalize_text("\n".join(buffer))
        if body:
            segments.append(Segment(text=body, section=section))
        buffer.clear()

    for line in text.splitlines():
        heading = re.match(r"^\s{0,3}(#{1,6})\s+(.*?)\s*#*\s*$", line)
        if heading:
            flush()
            section = heading.group(2).strip()
            continue
        line = re.sub(r"!\[([^\]]*)\]\([^)]*\)", r"\1", line)
        line = re.sub(r"\[([^\]]+)\]\([^)]*\)", r"\1", line)
        line = re.sub(r"[*_`]{1,3}", "", line)
        buffer.append(line)
    flush()
    return segments


def guess_title(filename, segments):
    for seg in segments[:3]:
        if seg.section and 3 <= len(seg.section) <= 120:
            return seg.section
    base = re.sub(r"\.[a-z0-9]+$", "", filename, flags=re.I)
    return re.sub(r"[_\-]+", " ", base).strip() or filename
