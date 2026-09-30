import io

import docx

from tests.conftest import register, upload


def test_course_crud_and_isolation(client, auth, course):
    assert course["code"] == "CS2028"
    dup = client.post("/api/courses", json={"name": "Again", "code": "CS2028"}, headers=auth)
    assert dup.status_code == 409

    resp = client.patch(f"/api/courses/{course['id']}", json={"exam_date": "2026-12-10"}, headers=auth)
    assert resp.get_json()["course"]["exam_date"] == "2026-12-10"

    other = register(client, email="other@example.edu")
    other_auth = {"Authorization": f"Bearer {other['access_token']}"}
    assert client.get(f"/api/courses/{course['id']}", headers=other_auth).status_code == 404

    listing = client.get("/api/courses", headers=auth).get_json()
    assert listing["meta"]["total"] == 1


def test_upload_markdown_is_chunked(client, auth, ready_doc):
    assert ready_doc["chunk_count"] >= 3
    assert ready_doc["title"]
    detail = client.get(f"/api/documents/{ready_doc['id']}", headers=auth).get_json()
    assert "Heaps and Priority Queues" in detail["sections"]
    chunks = client.get(f"/api/documents/{ready_doc['id']}/chunks?per_page=100", headers=auth).get_json()
    assert chunks["meta"]["total"] == ready_doc["chunk_count"]
    assert all(c["text"] for c in chunks["items"])


def test_duplicate_upload_is_detected(client, auth, course, ready_doc):
    again = upload(client, auth, course["id"])
    assert again.status_code == 200
    assert again.get_json()["duplicate"] is True
    assert again.get_json()["document"]["id"] == ready_doc["id"]


def test_bad_files_rejected(client, auth, course):
    fake_pdf = client.post(
        f"/api/courses/{course['id']}/documents",
        data={"file": (io.BytesIO(b"hello, not a pdf"), "notes.pdf")},
        headers=auth, content_type="multipart/form-data",
    )
    assert fake_pdf.status_code == 422
    exe = client.post(
        f"/api/courses/{course['id']}/documents",
        data={"file": (io.BytesIO(b"MZ..."), "virus.exe")},
        headers=auth, content_type="multipart/form-data",
    )
    assert exe.status_code == 422
    missing = client.post(f"/api/courses/{course['id']}/documents", headers=auth)
    assert missing.status_code == 422


def test_docx_upload(client, auth, course):
    document = docx.Document()
    document.add_heading("Graph Traversal", level=1)
    document.add_paragraph(
        "Breadth first search explores a graph level by level using a queue. It finds shortest paths in "
        "unweighted graphs. Depth first search uses a stack or recursion and is the basis for topological sort."
    )
    buf = io.BytesIO()
    document.save(buf)
    resp = client.post(
        f"/api/courses/{course['id']}/documents",
        data={"file": (io.BytesIO(buf.getvalue()), "graphs.docx")},
        headers=auth, content_type="multipart/form-data",
    )
    doc = resp.get_json()["document"]
    assert doc["status"] == "ready", doc
    assert doc["title"] == "Graph Traversal"


def test_empty_text_file_fails_cleanly(client, auth, course):
    resp = upload(client, auth, course["id"], text="   \n\n  ", filename="blank.txt")
    doc = resp.get_json()["document"]
    assert doc["status"] == "failed"
    assert "No readable text" in doc["error"]


def test_delete_and_reprocess(client, auth, ready_doc):
    resp = client.post(f"/api/documents/{ready_doc['id']}/reprocess", headers=auth)
    assert resp.status_code == 202
    assert resp.get_json()["document"]["version"] == 2
    assert client.delete(f"/api/documents/{ready_doc['id']}", headers=auth).status_code == 204
    assert client.get(f"/api/documents/{ready_doc['id']}", headers=auth).status_code == 404


def test_download_roundtrip(client, auth, ready_doc):
    resp = client.get(f"/api/documents/{ready_doc['id']}/download", headers=auth)
    assert resp.status_code == 200
    assert b"Binary Search Trees" in resp.data


def test_summaries_are_cached(client, auth, ready_doc):
    first = client.get(f"/api/documents/{ready_doc['id']}/summary?style=bullets", headers=auth).get_json()
    assert first["cached"] is False
    assert first["summary"]["content"].startswith("- ")
    second = client.get(f"/api/documents/{ready_doc['id']}/summary?style=bullets", headers=auth).get_json()
    assert second["cached"] is True
    bad = client.get(f"/api/documents/{ready_doc['id']}/summary?style=poem", headers=auth)
    assert bad.status_code == 422
